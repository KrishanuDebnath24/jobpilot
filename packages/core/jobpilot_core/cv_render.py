"""Deterministic CV rendering (Phase 6.5, 6.5e) — Irish-format, ATS-safe. NO LLM.

Renders a tailored structured CV to **PDF** and **DOCX** honouring Irish conventions and
ATS-parseability:

* **A4**, single-column, standard fonts, **selectable text** (never text-in-images or
  layout tables that break parsing);
* **no photo / date of birth / gender / marital status** (Irish anti-discrimination
  norm) — the renderer simply never emits them;
* Irish section order: **Profile → Key Skills → Experience → Education → "References
  available on request"**; British/Irish spelling is the content's responsibility.

PDF uses **reportlab** (pure-Python, no system libraries — so it renders identically on
Windows/CI and in the Linux image, and the machine-readability check runs everywhere).
WeasyPrint is the work-order's named tool but needs GTK/pango/cairo system libs that
aren't available on Windows/CI; reportlab meets every constraint above and is verifiable
in tests. DOCX uses **python-docx**. ``pdf_text``/``docx_text`` extract the rendered text
so callers can prove the output is machine-readable.
"""

from __future__ import annotations

import io

REFERENCES_LINE = "References available on request."


def _lines(cv: dict) -> tuple[dict, str, list[str], list[dict], list[str]]:
    contact = cv.get("contact") or {}
    summary = (cv.get("summary") or "").strip()
    skills = [s for s in (cv.get("skills") or []) if s and str(s).strip()]
    experience = cv.get("experience") or []
    education = [e.get("text", "").strip() if isinstance(e, dict) else str(e).strip()
                 for e in (cv.get("education") or [])]
    return contact, summary, skills, experience, [e for e in education if e]


def _contact_line(contact: dict) -> str:
    bits = [contact.get("email"), contact.get("phone"), contact.get("location")]
    bits += list(contact.get("links") or [])
    return "  •  ".join(b for b in bits if b)


# --------------------------------------------------------------------------- #
# PDF (reportlab) — A4, single column, Helvetica, selectable text
# --------------------------------------------------------------------------- #
def render_pdf(cv: dict) -> bytes:
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        HRFlowable, ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer,
    )

    def esc(s: str) -> str:
        return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    contact, summary, skills, experience, education = _lines(cv)

    name_st = ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=18, leading=22, alignment=TA_LEFT)
    meta_st = ParagraphStyle("meta", fontName="Helvetica", fontSize=9.5, leading=13, textColor="#333333")
    h_st = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=12, leading=16, spaceBefore=10, spaceAfter=2)
    body_st = ParagraphStyle("body", fontName="Helvetica", fontSize=10, leading=14)
    role_st = ParagraphStyle("role", fontName="Helvetica-Bold", fontSize=10.5, leading=14, spaceBefore=6)
    dates_st = ParagraphStyle("dates", fontName="Helvetica-Oblique", fontSize=9, leading=12, textColor="#555555")

    story: list = []
    story.append(Paragraph(esc(contact.get("name") or "Curriculum Vitae"), name_st))
    cl = _contact_line(contact)
    if cl:
        story.append(Paragraph(esc(cl), meta_st))
    story.append(Spacer(1, 4))
    story.append(HRFlowable(width="100%", thickness=0.6, color="#999999"))

    def section(title: str):
        story.append(Paragraph(esc(title), h_st))

    if summary:
        section("Profile")
        story.append(Paragraph(esc(summary), body_st))
    if skills:
        section("Key Skills")
        story.append(Paragraph(esc("  •  ".join(skills)), body_st))
    if experience:
        section("Experience")
        for e in experience:
            head = " — ".join(x for x in (e.get("title"), e.get("company")) if x)
            if head:
                story.append(Paragraph(esc(head), role_st))
            if e.get("dates"):
                story.append(Paragraph(esc(e["dates"]), dates_st))
            bullets = [b for b in (e.get("bullets") or []) if b and str(b).strip()]
            if bullets:
                story.append(ListFlowable(
                    [ListItem(Paragraph(esc(b), body_st), leftIndent=10) for b in bullets],
                    bulletType="bullet", start="•", leftIndent=12))
    if education:
        section("Education")
        for ed in education:
            story.append(Paragraph(esc(ed), body_st))
    story.append(Spacer(1, 8))
    story.append(Paragraph(esc(REFERENCES_LINE), dates_st))

    buf = io.BytesIO()
    SimpleDocTemplate(
        buf, pagesize=A4, topMargin=16 * mm, bottomMargin=16 * mm,
        leftMargin=18 * mm, rightMargin=18 * mm, title="CV",
    ).build(story)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# DOCX (python-docx) — A4, single column, selectable text
# --------------------------------------------------------------------------- #
def render_docx(cv: dict) -> bytes:
    import docx
    from docx.shared import Mm, Pt

    contact, summary, skills, experience, education = _lines(cv)
    document = docx.Document()
    for section in document.sections:  # A4
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(16)
        section.left_margin = section.right_margin = Mm(18)

    name = document.add_paragraph()
    run = name.add_run(contact.get("name") or "Curriculum Vitae")
    run.bold = True
    run.font.size = Pt(18)
    cl = _contact_line(contact)
    if cl:
        document.add_paragraph(cl)

    def heading(title: str):
        p = document.add_paragraph()
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(12)

    if summary:
        heading("Profile")
        document.add_paragraph(summary)
    if skills:
        heading("Key Skills")
        document.add_paragraph("  •  ".join(skills))
    if experience:
        heading("Experience")
        for e in experience:
            head = " — ".join(x for x in (e.get("title"), e.get("company")) if x)
            if head:
                p = document.add_paragraph()
                p.add_run(head).bold = True
            if e.get("dates"):
                document.add_paragraph(e["dates"]).italic = True
            for b in (e.get("bullets") or []):
                if b and str(b).strip():
                    document.add_paragraph(str(b), style="List Bullet")
    if education:
        heading("Education")
        for ed in education:
            document.add_paragraph(ed)
    document.add_paragraph(REFERENCES_LINE)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# Machine-readability checks — extract text back out of the rendered files
# --------------------------------------------------------------------------- #
def pdf_text(data: bytes) -> str:
    """Extract selectable text from a rendered PDF (proves it's not an image)."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages).strip()


def docx_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    return "\n".join(p.text for p in document.paragraphs).strip()


def render_both(cv: dict) -> dict:
    """Render PDF + DOCX; return bytes + a machine-readability verdict for each."""
    pdf, docx_bytes = render_pdf(cv), render_docx(cv)
    ptext = pdf_text(pdf)
    return {
        "pdf": pdf,
        "docx": docx_bytes,
        "pdf_text_len": len(ptext),
        "pdf_machine_readable": len(ptext) >= 50,  # real extractable text, not an image
    }
